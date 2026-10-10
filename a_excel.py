"""Subdivide columnas compuestas y exporta a Excel.

Lee salida/puestos_enriquecidos.csv y escribe salida/puestos.xlsx:
  ubicacion "Prov / Cantón / Distrito" -> provincia, canton, distrito
  fecha_vencimiento "miércoles, 30 de septiembre del 2026" -> fecha real de Excel
  nivel_academico "... (Requerido)"    -> nivel_academico + nivel_requerido

Opcionales:
  --vigentes   solo vacantes con fecha de vencimiento >= hoy (no vencidas)
  --nuevos     solo puestos NO registrados en salida/historial_ids.txt;
               destino por defecto salida/puestos_nuevos_<fecha>.xlsx; si hay
               nuevos, se fusionan con los ya guardados ese mismo dia y el
               historial se actualiza con todos los ids actuales; si no hay
               nuevos, no se toca el xlsx ni el historial
  --salida R   ruta del xlsx (default: salida/puestos.xlsx)
"""
import argparse
import csv
import datetime as dt
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("Falta openpyxl: .venv\\Scripts\\python.exe -m pip install openpyxl")

SRC = "salida/puestos_enriquecidos.csv"
DST = "salida/puestos.xlsx"
HISTORIAL = "salida/historial_ids.txt"

MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
         "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
         "noviembre": 11, "diciembre": 12}
RE_FECHA = re.compile(
    r"^\s*[a-záéíóúñ]+,?\s+(\d{1,2})\s+de\s+([a-záéíóúñ]+)\s+del?\s+(\d{4})\s*$", re.I)
RE_REQ = re.compile(r"\s*\((?:Requerido|Deseable|Opcional)\)\s*$", re.I)
RE_TIENE_REQ = re.compile(r"\((?:Requerido|Deseable|Opcional)\)\s*$", re.I)

ANCHO_MAX = 55
ANCHOS_FIJOS = {"url": 40, "correo": 30, "descripcion": 55, "requisitos": 55,
                "habilidades": 45, "titulo": 45, "ocupacion": 30}
NUMERICOS_INT = ("id", "plazas", "salario_min_crc", "salario_max_crc")
NUMERICOS_FLOAT = ("experiencia_min_anios",)


def a_numero(valor, entero):
    v = (valor or "").strip()
    if not v:
        return None
    try:
        return int(float(v)) if entero else float(v)
    except ValueError:
        return valor


def parsear_fecha(txt):
    m = RE_FECHA.match(txt or "")
    if not m:
        return None
    dia, mes, anio = int(m.group(1)), m.group(2).lower(), int(m.group(3))
    if mes not in MESES:
        return None
    try:
        return dt.date(anio, MESES[mes], dia)
    except ValueError:
        return None


def parsear_requerido(txt):
    v = (txt or "").strip()
    return (RE_REQ.sub("", v).strip(), "si" if RE_TIENE_REQ.search(v) else "no")


def a_id(f):
    try:
        return int(f["id"])
    except (KeyError, TypeError, ValueError):
        return None


def leer_historial():
    """Devuelve (conjunto de ids, maximo id) de HISTORIAL; (set(), 0) si no existe."""
    try:
        with open(HISTORIAL, encoding="utf-8") as fh:
            ids = {int(line.strip()) for line in fh if line.strip().isdigit()}
    except FileNotFoundError:
        return set(), 0
    return ids, (max(ids) if ids else 0)


def guardar_historial(ids):
    with open(HISTORIAL, "w", encoding="utf-8", newline="") as fh:
        fh.write("".join(f"{i}\n" for i in sorted(ids)))


def filas_previas(dst):
    """Filas de datos de un xlsx ya generado (p.ej. mismo dia); [] si no existe."""
    try:
        wb = load_workbook(dst, read_only=True)
    except FileNotFoundError:
        return []
    try:
        ws = wb["puestos"]
        return [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    except (KeyError, StopIteration):
        return []
    finally:
        wb.close()


def construir_encabezado(campos):
    out = []
    for c in campos:
        if c == "ubicacion":
            out += ["provincia", "canton", "distrito"]
        elif c == "nivel_academico":
            out += ["nivel_academico", "nivel_requerido"]
        else:
            out.append(c)
    return out


def convertir(f):
    fila = []
    for c in CAMPOS:
        if c == "ubicacion":
            partes = [p.strip() for p in (f.get("ubicacion") or "").split("/")]
            while len(partes) < 3:
                partes.append("")
            fila += partes[:3]
        elif c == "nivel_academico":
            nivel, req = parsear_requerido(f.get("nivel_academico", ""))
            fila += [nivel, req]
        elif c == "fecha_vencimiento":
            fecha = parsear_fecha(f.get("fecha_vencimiento", ""))
            fila.append(fecha if fecha else (f.get("fecha_vencimiento") or ""))
            ESTADO["fechas_no_parseadas"] += (fecha is None)
        elif c in NUMERICOS_INT:
            fila.append(a_numero(f.get(c, ""), True))
        elif c in NUMERICOS_FLOAT:
            fila.append(a_numero(f.get(c, ""), False))
        else:
            fila.append(f.get(c, ""))
    return fila


def main():
    ap = argparse.ArgumentParser(description="Subdivide columnas y exporta a Excel")
    ap.add_argument("--vigentes", action="store_true",
                    help="solo vacantes con fecha de vencimiento >= hoy")
    ap.add_argument("--nuevos", action="store_true",
                    help="solo puestos no registrados en el historial (sin nuevos "
                         "no escribe el xlsx ni toca el historial)")
    ap.add_argument("--salida", default=None,
                    help="ruta del xlsx de salida (default: salida/puestos.xlsx; "
                         "con --nuevos: salida/puestos_nuevos_<fecha>.xlsx)")
    args = ap.parse_args()

    if args.salida:
        dst = args.salida
    elif args.nuevos:
        dst = f"salida/puestos_nuevos_{dt.date.today().isoformat()}.xlsx"
    else:
        dst = DST

    with open(SRC, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        raise SystemExit(f"{SRC} vacio o no existe (correr reglas.py)")

    global CAMPOS, ESTADO
    CAMPOS = list(rows[0].keys())
    ESTADO = {"fechas_no_parseadas": 0}

    ids_actuales = set()
    for f in rows:
        i = a_id(f)
        if i is not None:
            ids_actuales.add(i)

    hist_antes, hist_max = set(), 0
    if args.nuevos:
        hist_antes, hist_max = leer_historial()
        if hist_antes:
            rows = [f for f in rows
                    if (i := a_id(f)) is not None
                    and i not in hist_antes]
        print(f"historial: {len(hist_antes)} ids, ultimo = {hist_max}"
              if hist_antes else "historial vacio (primera corrida): se exportan todos")
        print(f"nuevos: {len(rows)} de {len(ids_actuales)} puestos actuales")

    encabezado = construir_encabezado(CAMPOS)
    datos = [convertir(f) for f in rows]
    total = len(datos)

    excluidos = sin_fecha = 0
    if args.vigentes:
        i_fec = encabezado.index("fecha_vencimiento")
        hoy = dt.date.today()
        conservadas = []
        for fila in datos:
            v = fila[i_fec]
            if isinstance(v, dt.date):
                if v >= hoy:
                    conservadas.append(fila)
                else:
                    excluidos += 1
            else:
                conservadas.append(fila)
                sin_fecha += 1
        datos = conservadas

    if args.nuevos and not datos:
        print("  sin puestos nuevos: se conserva el xlsx anterior y el historial")
        return

    fusion = 0
    if args.nuevos and not args.salida and datos:
        previas = filas_previas(dst)
        if previas and len(previas[0]) != len(encabezado):
            previas = []
        if previas:
            i_id = encabezado.index("id")
            ya = {f[i_id] for f in previas}
            extra = [f for f in datos if f[i_id] not in ya]
            fusion = len(extra)
            datos = previas + extra

    wb = Workbook()
    ws = wb.active
    ws.title = "puestos"
    ws.append(encabezado)
    for fila in datos:
        ws.append(fila)

    hdr_fill = PatternFill("solid", fgColor="D9E1F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = hdr_fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(encabezado))}{len(datos) + 1}"

    for idx, nombre in enumerate(encabezado, start=1):
        if nombre in ANCHOS_FIJOS:
            ancho = ANCHOS_FIJOS[nombre]
        else:
            max_len = max((len(str(ws.cell(row=r, column=idx).value or ""))
                           for r in range(1, min(len(datos), 300) + 1)), default=10)
            ancho = min(max(max_len + 2, 10), 24)
        ws.column_dimensions[get_column_letter(idx)].width = ancho

    for fila in ws.iter_rows(min_row=2):
        for cell in fila:
            if isinstance(cell.value, dt.date):
                cell.number_format = "yyyy-mm-dd"

    try:
        wb.save(dst)
    except PermissionError:
        raise SystemExit(
            f"ERROR: no se pudo escribir {dst}\n"
            "  Causa: el archivo esta ABIERTO en Excel (bloqueado).\n"
            "  Cierre ese archivo en Excel y vuelva a ejecutar.")

    if args.nuevos:
        guardar_historial(hist_antes | ids_actuales)
        print(f"  historial actualizado: {len(hist_antes)} -> "
              f"{len(hist_antes | ids_actuales)} ids -> {HISTORIAL}")
        if fusion:
            print(f"  fusion con archivo previo del dia: +{fusion} filas nuevas")

    check = load_workbook(dst)
    ws2 = check["puestos"]
    n_cols = ws2.max_column
    n_filas = ws2.max_row
    print(f"OK: {dst} | {n_filas - 1} filas x {n_cols} columnas")
    if args.vigentes:
        print(f"  filtro vigentes: {total} -> {total - excluidos} "
              f"(excluidos vencidos: {excluidos}, sin fecha parseada: {sin_fecha})")
    print(f"  fechas vencimiento no parseadas: {ESTADO['fechas_no_parseadas']}")
    hdr = [c.value for c in ws2[1]]
    print("  encabezado:", ",".join(hdr))
    muestra = next((r for r in ws2.iter_rows(min_row=2, values_only=True)
                    if r[hdr.index("provincia")]), None)
    if muestra is None:
        return
    for campo in ("provincia", "canton", "distrito", "fecha_vencimiento",
                  "nivel_academico", "nivel_requerido"):
        print(f"  ejemplo {campo}: {muestra[hdr.index(campo)]!r}")


if __name__ == "__main__":
    main()
